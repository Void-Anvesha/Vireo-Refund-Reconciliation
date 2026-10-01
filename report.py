"""Board pack, chart and generated memo from reconciled numbers."""
import json
import math
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from load_clean import ROOT
from policy_checks import FLAG_COLUMNS


def summary_table(frame, keys):
    return frame.groupby(keys, dropna=False, sort=True).agg(refund_count=('ticket_id', 'size'), refund_minor=('refund_minor', 'sum')).reset_index()


def insights(refunds, analysis, costs):
    total = int(refunds.refund_minor.sum())
    change = analysis['decomposition'].iloc[0]
    latest = analysis['opportunity'].iloc[-1]
    gw = refunds[refunds.quarter.eq(latest.quarter) & refunds.goodwill_original_over_cap].refund_minor.sum()
    eval_path = ROOT / 'output/evaluation/metrics.json'
    if eval_path.exists():
        try:
            eval_metrics = json.loads(eval_path.read_text(encoding='utf-8'))
            if eval_metrics.get('labeled_tickets'):
                accuracy_note = f"Diagnostic rules accuracy: {eval_metrics['rules']['accuracy']:.1%} on {eval_metrics['labeled_tickets']} reviewed tickets."
            else:
                accuracy_note = "Classification accuracy is not yet human-validated."
        except (KeyError, ValueError, json.JSONDecodeError):
            accuracy_note = "Classification accuracy is not yet human-validated."
    else:
        accuracy_note = "Classification accuracy is not yet human-validated."
    return [f"Total refunds reconcile to Rs {total / 1e7:,.2f} lakh across {len(refunds):,} tickets.",
            f"Q2 2025 to Q4 2025 rose Rs {change.change_inr / 1e5:,.2f} lakh: volume contributed Rs {change.volume_effect_inr / 1e5:,.2f} lakh and average value Rs {change.average_effect_inr / 1e5:,.2f} lakh.",
            f"Original-coded goodwill above Rs 500 totalled Rs {gw / 1e7:,.2f} lakh in {latest.quarter}; reason review is required before treating this as a goodwill breach.",
            f"{int(refunds.refund_and_replacement.sum())} tickets have both refund and replacement; prioritize an order-level check before compensation.",
            f"Classification mode: {costs['mode']}; {int(refunds.needs_review.sum()):,} tickets need review. {accuracy_note}"]


def present(frame):
    frame = frame.copy()
    for col in list(frame.columns):
        if col.endswith('_minor'):
            frame[col] = frame[col] / 1e7
            frame = frame.rename(columns={col: col[:-6] + ' (Rs lakh)'})
        elif col.endswith('_inr'):
            frame[col] = frame[col] / 1e5
            frame = frame.rename(columns={col: col[:-4] + ' (Rs lakh)'})
    return frame


def safe_cell(value):
    if pd.isna(value):
        return None
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    if hasattr(value, 'item'):
        return value.item()
    return value


def write_table(ws, title, frame, row=1):
    ws.cell(row, 1, title).font = Font(bold=True, size=13, color='16324F')
    frame = present(frame)
    for j, col in enumerate(frame.columns, 1):
        cell = ws.cell(row + 1, j, str(col).replace('_', ' '))
        cell.fill = PatternFill('solid', fgColor='16324F')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for i, values in enumerate(frame.itertuples(index=False, name=None), row + 2):
        for j, value in enumerate(values, 1):
            cell = ws.cell(i, j, safe_cell(value))
            name = str(frame.columns[j - 1])
            if 'Rs lakh' in name:
                cell.number_format = '#,##0.000;[Red](#,##0.000);-'
            elif 'rate' in name or 'share' in name:
                cell.number_format = '0.0%'
            elif 'confidence' in name or 'csat' in name:
                cell.number_format = '0.00'
            elif isinstance(value, (float, int)):
                cell.number_format = '#,##0.##'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if i % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='EDF3F8')
    return row + len(frame) + 4


def build_report(all_tickets, refunds, waterfall, audit, analysis, costs):
    output = ROOT / 'output'
    total = int(refunds.refund_minor.sum())
    monthly = summary_table(refunds, ['month'])
    team = summary_table(refunds, ['month', 'team', 'tier'])
    reasons = summary_table(refunds, ['month', 'final_reason'])
    agent = refunds.groupby(['month', 'team', 'tier', 'agent_id', 'agent_name', 'site'], dropna=False).agg(
        refund_count=('ticket_id', 'size'), refund_minor=('refund_minor', 'sum'),
        flagged_tickets=('policy_flag_count', lambda s: int(s.gt(0).sum())), **{x: (x, 'sum') for x in FLAG_COLUMNS}).reset_index()
    volumes = refunds.groupby('agent_id').size()
    agent['period_refund_count'] = agent.agent_id.map(volumes)
    agent['volume_flag'] = agent.period_refund_count.map(lambda n: 'low volume' if n < 20 else '20+ refunds')
    mismatch = summary_table(refunds, ['refund_reason_code', 'inferred_reason'])
    status = summary_table(refunds, ['status'])
    coverage = []
    for name, table in [('Summary', monthly), ('By Team', team), ('By Reason', reasons), ('By Agent', agent), ('Mismatch Matrix', mismatch), ('Status', status), ('Quarters', analysis['quarters'])]:
        assert int(table.refund_minor.sum()) == total, f'{name} does not reconcile'
        assert int(table.refund_count.sum()) == len(refunds), f'{name} duplicates or loses tickets'
        coverage.append({'table': name, 'refund_minor': int(table.refund_minor.sum()), 'scope': 'Full refund population; exact integer-paise reconciliation'})
    flagged = refunds[refunds.policy_flag_count.gt(0)]
    review = refunds[refunds.needs_review].sort_values('refund_minor', ascending=False)
    for name, subset in [('Policy Flags', flagged), ('Needs Review', review)]:
        coverage.append({'table': name, 'refund_minor': int(subset.refund_minor.sum()), 'scope': f'Filtered subset of {len(subset)} tickets, not additive to full total'})
        complement = refunds[~refunds.ticket_id.isin(subset.ticket_id)]
        assert int(subset.refund_minor.sum() + complement.refund_minor.sum()) == total
    fig, ax = plt.subplots(figsize=(11, 4.2))
    ax.plot(monthly.month, monthly.refund_minor / 1e7, color='#247B91', marker='o', linewidth=2.4)
    q4_positions = [i for i, month in enumerate(monthly.month) if month.startswith(('2025-10', '2025-11', '2025-12'))]
    if q4_positions:
        ax.axvspan(min(q4_positions) - .4, max(q4_positions) + .4, alpha=.13, color='#E89B3C', label='Q4 2025')
    ax.set(title='Monthly refunds | created month, IST', ylabel='Rs lakh', xlabel='')
    ax.tick_params(axis='x', rotation=45)
    ax.spines[['top', 'right']].set_visible(False)
    ax.grid(axis='y', alpha=.2)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(output / 'monthly_refunds.png', dpi=160)
    plt.close(fig)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    sentences = insights(refunds, analysis, costs)
    for i, line in enumerate(sentences, 1):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=9)
        ws.cell(i, 1, line).alignment = Alignment(wrap_text=True, vertical='center')
        ws.cell(i, 1).font = Font(bold=i == 1, color='16324F', size=12)
        ws.row_dimensions[i].height = 32
    row = write_table(ws, 'Monthly reconciled refunds | Rs lakh', monthly, 7)
    row = write_table(ws, 'Refund status split | open/pending included', status, row)
    write_table(ws, 'Separate costs | planning estimates, not part of refund total', analysis['separate_costs'], row)
    ws.freeze_panes = 'B9'
    chart = wb.create_sheet('Trend chart')
    chart['A1'] = 'Reconciled monthly refunds; shaded area is Q4 2025.'
    img = Image(str(output / 'monthly_refunds.png'))
    img.width, img.height = 1056, 403
    chart.add_image(img, 'A3')
    for name, title, table in [('By Team', 'Resolving team first | Returns Desk processes refunds by design; compare within tier', team),
                               ('By Reason', 'Final reasons | provisional until human evaluation', reasons),
                               ('By Agent', 'Alphabetical team and ID view; no performance ranking; low volume = <20 period refunds', agent),
                               ('Mismatch Matrix', 'Original agent code versus inferred code | all confidence levels', mismatch)]:
        ws = wb.create_sheet(name)
        write_table(ws, title, table)
        ws.auto_filter.ref = f'A2:{get_column_letter(len(table.columns))}{len(table) + 2}'
    ws = wb.create_sheet('Why Refunds Rose')
    row = 1
    for key, title in [('quarters', 'Quarter totals'), ('decomposition', 'Exact symmetric decomposition: count effect + average-value effect = change'),
                       ('claim', 'Priya claim: goodwill-like = GW-OTHER + PRICE-ADJ; classifications provisional'),
                       ('csat', 'CSAT: completed tickets only; blanks excluded; nonresponse and case mix limit causality'),
                       ('opportunity', 'Below 1% target: avoided count x observed mean replacement/refund value; benefits are alternatives, not additive'),
                       ('contributions', 'Contribution to change by reason, resolving team and channel')]:
        row = write_table(ws, title, analysis[key], row)
    detail_columns = ['ticket_id', 'month', 'team', 'tier', 'agent_id', 'refund_minor', 'refund_reason_code', 'inferred_reason', 'final_reason', 'confidence', 'classification_source', 'order_match', 'policy_flags', 'evidence']
    write_table(wb.create_sheet('Policy Flags'), 'Screening flags, not confirmed breaches; resolving agent may not be approving agent', flagged[detail_columns])
    review_columns = detail_columns + ['customer_message', 'agent_notes']
    write_table(wb.create_sheet('Needs Review'), 'Below 0.7 confidence | largest refund first | original code retained', review[review_columns])
    ws = wb.create_sheet('Reconciliation')
    row = write_table(ws, 'Mixed-unit export to reconciled refund spend', waterfall)
    row = write_table(ws, 'Table coverage | filtered sheets and comparisons are not additive', pd.DataFrame(coverage), row)
    row = write_table(ws, 'Order matching audit', refunds.order_match.value_counts().rename_axis('match_method').reset_index(name='refund_tickets'), row)
    audit_frame = pd.DataFrame([{'metric': k, 'value': json.dumps(v) if isinstance(v, dict) else str(v)} for k, v in audit.items()])
    write_table(ws, 'Load audit', audit_frame, row)
    decisions = [line[2:] for line in (ROOT / 'DECISIONS.md').read_text(encoding='utf-8').splitlines() if line.startswith('- ')]
    ws = wb.create_sheet('Assumptions')
    write_table(ws, 'Decisions and interpretation', pd.DataFrame({'decision': decisions}))
    ws.column_dimensions['A'].width = 115
    for ws in wb:
        if ws.freeze_panes is None:
            ws.freeze_panes = 'A3'
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '1:2' if ws.title != 'Summary' else '1:8'
        if ws.title != 'Assumptions':
            for column in ws.columns:
                letter = get_column_letter(column[0].column)
                values = [len(str(cell.value or '')) for cell in list(column)[1:80]]
                ws.column_dimensions[letter].width = min(48, max(16, max(values, default=16) + 2))
        for cells in ws.iter_rows():
            if ws.title == 'Summary' and cells[0].row <= 5:
                continue
            lines = max((math.ceil(len(str(c.value or '')) / max(10, ws.column_dimensions[get_column_letter(c.column)].width - 2)) for c in cells if c.value is not None), default=1)
            ws.row_dimensions[cells[0].row].height = min(120, max(24, lines * 15))
        ws.sheet_properties.outlinePr.summaryRight = False
    try:
        wb.save(output / 'vireo_refunds.xlsx')
    except PermissionError as exc:
        raise PermissionError(
            'Cannot overwrite output/vireo_refunds.xlsx. Close it in Excel or any preview pane, then run again.'
        ) from exc
    for key, frame in analysis.items():
        frame.to_csv(output / f'{key}.csv', index=False)
    pd.DataFrame(coverage).to_csv(output / 'table_reconciliation.csv', index=False)
    waterfall.to_csv(output / 'reconciliation.csv', index=False)
    export_cols = ['ticket_id', 'refund_reason_code', 'inferred_reason', 'final_reason', 'confidence', 'evidence', 'classification_source', 'rules_reason', 'rules_confidence', 'refund_minor', 'channel', 'source_system']
    refunds[export_cols].to_csv(output / 'classifications.csv', index=False)
    (output / 'audit.json').write_text(json.dumps(audit, indent=2), encoding='utf-8')
    write_memo(refunds, analysis, costs, total)
    return sentences


LEFT_OUT = ('Interactive dashboard: a static board pack meets the request. Live refund guardrail: recommendation only. '
            'Lot/batch drill-down: missing and ambiguous orders weaken coverage. Agent fairness model and ranking: case mix and small samples make ranking unfair. '
            'Question-answer chat: unnecessary for this board pack. Repeat-customer, product and deeper channel drill-down: outside the core time budget.')


def write_memo(refunds, analysis, costs, total):
    q = analysis['quarters'].set_index('quarter')
    change = analysis['decomposition'].iloc[0]
    opportunity = analysis['opportunity'].iloc[-1]
    reasons = refunds.groupby('final_reason').refund_minor.sum().sort_values(ascending=False).head(3)
    teams = refunds.groupby('team').refund_minor.sum().sort_values(ascending=False).head(3)
    csat = analysis['csat'][analysis['csat'].population.eq('All tickets')].set_index('quarter')
    claim = analysis['claim'].set_index('quarter')
    a, b = claim.loc['2025Q2'], claim.loc['2025Q4']
    # One central number is explicitly the conservative replacement-cost estimate.
    eval_path = ROOT / 'output/evaluation/metrics.json'
    if eval_path.exists():
        try:
            eval_metrics = json.loads(eval_path.read_text(encoding='utf-8'))
            if eval_metrics.get('labeled_tickets'):
                eval_sentence = (
                    f"A small diagnostic review of {eval_metrics['labeled_tickets']} labeled tickets gives the rules baseline "
                    f"{eval_metrics['rules']['accuracy']:.1%} accuracy, so reason-code estimates should be treated as directional until more labels are reviewed. "
                )
            else:
                eval_sentence = "Human evaluation remains pending, so reason-code accuracy is not yet measured. "
        except (KeyError, ValueError, json.JSONDecodeError):
            eval_sentence = "Human evaluation remains pending, so reason-code accuracy is not yet measured. "
    else:
        eval_sentence = "Human evaluation remains pending, so reason-code accuracy is not yet measured. "
    api_sentence = "The paid AI pilot was not run. " if costs.get('calls', 0) == 0 else "The API pilot cost and results are recorded in output/run_cost.json. "
    memo = f'''# To Arjun Mehta - refund review

Refunds were Rs {q.loc['2025Q2', 'refund_minor']/1e7:.2f} lakh in Q2 2025, Rs {q.loc['2025Q4', 'refund_minor']/1e7:.2f} lakh in Q4 2025, and Rs {q.loc['2026Q2', 'refund_minor']/1e7:.2f} lakh in Q2 2026.
The inflated export mixed legacy paise with rupees and counted migration copies twice.
After corrections, the full period reconciles to **Rs {total/1e7:.2f} lakh** across {len(refunds):,} refunds; the workbook shows every adjustment.

**What and who.** The largest provisional reasons are {', '.join(f'{name} (Rs {value/1e7:.2f} lakh)' for name, value in reasons.items())}. Classification mode is **{costs['mode']}**: {int(refunds.needs_review.sum()):,} uncertain tickets retain their original code and need review. Leading processing teams are {', '.join(f'{name} (Rs {value/1e7:.2f} lakh)' for name, value in teams.items())}. Returns Desk processes refunds by design. The workbook lists agents within teams, flags low volumes and separates tiers; amounts are not a performance ranking or evidence of blame.

**Why refunds rose.** Q2 to Q4 2025 increased by Rs {change.change_inr/1e5:.2f} lakh: refund-count growth explains Rs {change.volume_effect_inr/1e5:.2f} lakh and higher average refunds Rs {change.average_effect_inr/1e5:.2f} lakh, using a symmetric decomposition. Goodwill-like ticket count changed from {int(a.goodwill_like_count)} to {int(b.goodwill_like_count)}, DOA from {int(a.doa_count)} to {int(b.doa_count)}, and transit from {int(a.transit_count)} to {int(b.transit_count)}. All-ticket CSAT moved from {csat.loc['2025Q2', 'mean_csat']:.2f} to {csat.loc['2025Q4', 'mean_csat']:.2f} (change {csat.loc['2025Q4', 'mean_csat']-csat.loc['2025Q2', 'mean_csat']:+.2f}); blanks are excluded. These comparisons test Priya's account, but provisional labels, survey nonresponse and changing case mix prevent a causal verdict. See the workbook for the Q2 2026 comparison and response counts.

**Money at stake.** Prioritize the refund-plus-replacement check. At Q2 2026 volumes, the conservative quarterly opportunity is **Rs {opportunity.saving_lower_inr/1e5:.2f} lakh**, with a replacement-cost-to-refund-value range of **Rs {opportunity.saving_lower_inr/1e5:.2f}-{opportunity.saving_upper_inr/1e5:.2f} lakh**. Arithmetic: {int(opportunity.double_tickets)} double-compensation tickets out of {int(opportunity.refund_tickets)} ({opportunity.double_rate:.1%}); fewer than 1% allows at most {int(opportunity.target_max_tickets)}, so avoid {int(opportunity.avoided_tickets)}. Multiply that count by the observed mean replacement cost or refund value. These are alternative avoided-cost scenarios, not amounts to add or guaranteed recoveries. At refund time, check the order's refund/replacement history across teams and require lead review of exceptions. Improve order capture so the check can work reliably.

**Reason-code target and uncertainty.** In Q2 2026 GW-OTHER changes from {opportunity.original_gw_share:.1%} to {opportunity.final_gw_share:.1%} of refund value after provisional recoding; target is below 15%. Recoding improves visibility, not cash savings. Double-compensation evidence is stronger because it uses explicit fields. {eval_sentence}{api_sentence}Resolving-agent attribution cannot prove who approved a replacement. Policy flags before April 2025 apply the supplied policy retrospectively. Ambiguous order matches remain unresolved.

**Left out.** {LEFT_OUT}
'''
    (ROOT / 'MEMO.md').write_text(memo, encoding='utf-8')
